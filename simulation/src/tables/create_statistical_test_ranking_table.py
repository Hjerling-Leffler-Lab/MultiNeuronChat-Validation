import os

import pickle

import numpy as np
import pandas as pd

from scipy import stats

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

import argparse

parser = argparse.ArgumentParser()
parser.add_argument(
    '--path_to_summary_pkls',
    type=str,
    required=True,
)
parser.add_argument(
    '--mean_type',
    type=str,
    required=True,
)
parser.add_argument(
    '--path_to_ranking_table_excel',
    type=str,
    required=True,
)
parser.add_argument(
    '--n_donors',
    type=int,
    default=80,
)
parser.add_argument(
    '--alpha',
    type=float,
    default=0.05,
)
args = parser.parse_args()

path_to_summary_pkls: str = args.path_to_summary_pkls
mean_type: str = args.mean_type
path_to_ranking_table_excel: str = args.path_to_ranking_table_excel
n_donors: int = args.n_donors
alpha: float = args.alpha

CASES: list[str] = ['CASE_0.3', 'CASE_0.5', 'CASE_1', 'CASE_1.5']
PROPORTIONS: list[str] = ['0.1', '0.3', '0.5', '0.8', '1.0']
STATISTICAL_TESTS: list[str] = ['KS', 'CVM', 'MannWhitneyU']
SHORT_LABELS: dict[str, str] = {'KS': 'KS', 'CVM': 'CvM', 'MannWhitneyU': 'MWU'}
PAIRS: list[tuple[str, str]] = [('KS', 'CVM'), ('KS', 'MannWhitneyU'), ('CVM', 'MannWhitneyU')]

# Layout of the original hand-made workbook: one block per Log2 FC case, 10 rows apart, starting at row 2.
BLOCK_START_ROWS: list[int] = [2, 12, 22, 32]
COLUMN_WIDTHS: dict[str, float] = {
    'B': 11.14, 'C': 2.71, 'E': 5.29, 'F': 5.71, 'G': 9.0, 'H': 10.0, 'I': 5.0, 'K': 7.29, 'L': 5.0,
    'O': 5.29, 'P': 12.0, 'R': 5.29, 'S': 12.0,
}
FONT: Font = Font(name='Aptos Narrow', size=11)
CENTER: Alignment = Alignment(horizontal='center')


def paired_wilcoxon_p_value(x: np.ndarray, y: np.ndarray) -> float:
    """
    Same test as the significance bars in plot_pr_curves_b.py: exact paired Wilcoxon signed-rank test.
    Identical PR AUCs (all differences zero) and NaN p-values count as p = 1.
    """
    if not np.any(x - y):
        return 1.0

    p_value: float = stats.wilcoxon(x, y, zero_method='wilcox', method=stats.PermutationMethod()).pvalue

    return 1.0 if np.isnan(p_value) else p_value


def pairwise_outcome(x: np.ndarray, y: np.ndarray, adjusted_p_value: float) -> int:
    """
    +1 if the first test wins, -1 if the second test wins, 0 if tied.

    A win needs a significant difference; the winner is the test with the higher median PR AUC (the boxplot
    median). If the medians are equal, the sign of the median paired difference decides.
    """
    if adjusted_p_value >= alpha:
        return 0

    median_difference: float = np.median(x) - np.median(y)
    if median_difference == 0:
        median_difference = np.median(x - y)

    return int(np.sign(median_difference))


def sheet_name_for_proportion(proportion: str) -> str:
    return f'{round(float(proportion) * n_donors)}_OutOf_{n_donors}'


def set_cell(ws: Worksheet, coordinate: str, value, center: bool = False):
    ws[coordinate] = value
    ws[coordinate].font = FONT
    if center:
        ws[coordinate].alignment = CENTER


def write_block(ws: Worksheet, start_row: int, log2_fc: str, outcomes: list[int]):
    """
    One Log2 FC block: pairwise outcomes (B:C), score and rank per test (E:I), tests sorted by rank (K:L).

    For three tests the score (net wins) ranges from -2 to 2. Pos Score shifts it to 1..5 and Rank Score maps it
    linearly onto 1 (score 2) .. 3 (score -2), the sort key for the Ranking column.
    """
    ws.merge_cells(f'B{start_row}:L{start_row}')
    set_cell(ws, f'B{start_row}', log2_fc, center=True)

    header_row: int = start_row + 1
    for column, header in zip('FGHIKL', ['Score', 'Pos Score', 'Rank Score', 'Rank', 'Ranking', 'Rank']):
        set_cell(ws, f'{column}{header_row}', header, center=True)

    first, second, third = start_row + 2, start_row + 3, start_row + 4

    for row, (a, b), outcome in zip([first, second, third], PAIRS, outcomes):
        set_cell(ws, f'B{row}', f'{SHORT_LABELS[a]} vs {SHORT_LABELS[b]}')
        set_cell(ws, f'C{row}', outcome)

    # Rows of E:I in the order of STATISTICAL_TESTS: KS, CvM, MWU.
    scores: list[str] = [
        f'=C{first}+C{second}',
        f'=(-C{first})+C{third}',
        f'=(-C{second})+(-C{third})',
    ]
    test_rows: list[int] = [first, second, third]
    for row, test, score in zip(test_rows, STATISTICAL_TESTS, scores):
        other_rows: list[int] = [other_row for other_row in test_rows if other_row != row]

        set_cell(ws, f'E{row}', SHORT_LABELS[test])
        set_cell(ws, f'F{row}', score)
        set_cell(ws, f'G{row}', f'=F{row}+3')
        set_cell(ws, f'H{row}', f'=ABS((F{row}-2)/4)*2 + 1')
        set_cell(ws, f'I{row}', f'=1+(G{row}<G{other_rows[0]}) +(G{row}<G{other_rows[1]})')

    ws[f'K{first}'] = ArrayFormula(f'K{first}:K{third}', f'=_xlfn.SORTBY(E{first}:E{third},H{first}:H{third})')
    ws[f'L{first}'] = ArrayFormula(f'L{first}:L{third}', f'=_xlfn.SORTBY(I{first}:I{third},H{first}:H{third})')


def write_mean_rank_summary(ws: Worksheet, title_row: int, title: str, mean_rank_formulas: list[str],
                            with_mean_rank_header: bool):
    ws.merge_cells(f'O{title_row}:S{title_row}')
    set_cell(ws, f'O{title_row}', title, center=True)

    ws.merge_cells(f'O{title_row + 1}:P{title_row + 1}')
    ws.merge_cells(f'R{title_row + 1}:S{title_row + 1}')
    set_cell(ws, f'O{title_row + 1}', 'Unsorted', center=True)
    set_cell(ws, f'R{title_row + 1}', 'Sorted', center=True)

    first: int = title_row + 2
    if with_mean_rank_header:
        set_cell(ws, f'P{first}', 'Mean Rank')
        set_cell(ws, f'S{first}', 'Mean Rank')
        first += 1
    last: int = first + len(STATISTICAL_TESTS) - 1

    for row, test, formula in zip(range(first, last + 1), STATISTICAL_TESTS, mean_rank_formulas):
        set_cell(ws, f'O{row}', SHORT_LABELS[test])
        set_cell(ws, f'P{row}', formula)

    ws[f'R{first}'] = ArrayFormula(f'R{first}:R{last}', f'=_xlfn.SORTBY(O{first}:O{last},P{first}:P{last})')
    ws[f'S{first}'] = ArrayFormula(f'S{first}:S{last}', f'=_xlfn.SORTBY(P{first}:P{last},P{first}:P{last})')


def main():
    pairwise_rows: list[dict] = []
    outcomes: dict[tuple[str, str], list[int]] = {}

    for proportion in PROPORTIONS:
        for case in CASES:
            path_to_pkl: str = os.path.join(path_to_summary_pkls, f'{mean_type}_{case}_{proportion}.pkl')

            if not os.path.exists(path_to_pkl):
                raise FileNotFoundError(f"File not found: {path_to_pkl}")

            with open(path_to_pkl, 'rb') as f:
                pr_aucs: dict[str, np.ndarray] = {
                    test: np.asarray(values) for test, values in pickle.load(f).items()
                }

            p_values: list[float] = [paired_wilcoxon_p_value(pr_aucs[a], pr_aucs[b]) for a, b in PAIRS]
            adjusted_p_values: np.ndarray = stats.false_discovery_control(p_values, method='bh')

            outcomes[(proportion, case)] = []
            for (a, b), p_value, adjusted_p_value in zip(PAIRS, p_values, adjusted_p_values):
                outcome: int = pairwise_outcome(pr_aucs[a], pr_aucs[b], adjusted_p_value)
                outcomes[(proportion, case)].append(outcome)

                pairwise_rows.append({
                    'Proportion of Affected Donors': proportion,
                    'Log2 FC': case.split('_')[-1],
                    'Comparison': f'{SHORT_LABELS[a]} vs {SHORT_LABELS[b]}',
                    'Median PR AUC – First': np.median(pr_aucs[a]),
                    'Median PR AUC – Second': np.median(pr_aucs[b]),
                    'Wilcoxon p': p_value,
                    'BH-adjusted p': adjusted_p_value,
                    'Outcome': outcome,
                })

    workbook: Workbook = Workbook()
    workbook.remove(workbook.active)

    for proportion in PROPORTIONS:
        ws: Worksheet = workbook.create_sheet(sheet_name_for_proportion(proportion))

        for column, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[column].width = width

        for start_row, case in zip(BLOCK_START_ROWS, CASES):
            write_block(ws, start_row, case.split('_')[-1], outcomes[(proportion, case)])

        # Mean rank of each test over the Log2 FC blocks (Rank column I, test rows 2..4 below each block start).
        write_mean_rank_summary(
            ws,
            title_row=3,
            title=f'Proportion {proportion}',
            mean_rank_formulas=[
                '=AVERAGE(' + ','.join(f'I{start_row + 2 + index}' for start_row in BLOCK_START_ROWS) + ')'
                for index in range(len(STATISTICAL_TESTS))
            ],
            with_mean_rank_header=True,
        )

    # Mean rank over all proportions, on the last sheet as in the original workbook.
    first_mean_rank_row: int = 6
    write_mean_rank_summary(
        workbook[sheet_name_for_proportion(PROPORTIONS[-1])],
        title_row=11,
        title='Over all proportions',
        mean_rank_formulas=[
            '=AVERAGE(' + ','.join(
                f"'{sheet_name_for_proportion(proportion)}'!P{first_mean_rank_row + index}"
                for proportion in PROPORTIONS
            ) + ')'
            for index in range(len(STATISTICAL_TESTS))
        ],
        with_mean_rank_header=False,
    )

    # The p-values and medians behind every +1/0/-1 in column C.
    pairwise_ws: Worksheet = workbook.create_sheet('Pairwise comparisons')
    pairwise_df: pd.DataFrame = pd.DataFrame(pairwise_rows)
    pairwise_ws.append(list(pairwise_df.columns))
    for row in pairwise_df.itertuples(index=False):
        pairwise_ws.append(list(row))

    workbook.calculation.fullCalcOnLoad = True
    workbook.save(path_to_ranking_table_excel)


if __name__ == "__main__":
    main()
