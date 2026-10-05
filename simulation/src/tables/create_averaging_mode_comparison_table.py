import os

import pickle

import numpy as np
import pandas as pd

from scipy import stats

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

import argparse

parser = argparse.ArgumentParser()
parser.add_argument(
    '--path_to_summary_pkls',
    type=str,
    required=True,
)
parser.add_argument(
    '--path_to_comparison_table_excel',
    type=str,
    required=True,
)
parser.add_argument(
    '--n_donors',
    type=int,
    default=80,
)
parser.add_argument(
    '--alpha',
    type=float,
    default=0.05,
)
args = parser.parse_args()

path_to_summary_pkls: str = args.path_to_summary_pkls
path_to_comparison_table_excel: str = args.path_to_comparison_table_excel
n_donors: int = args.n_donors
alpha: float = args.alpha

CASES: list[str] = ['CASE_0.3', 'CASE_0.5', 'CASE_1', 'CASE_1.5']
PROPORTIONS: list[str] = ['0.1', '0.3', '0.5', '0.8', '1.0']
STATISTICAL_TESTS: list[str] = ['KS', 'CVM', 'MannWhitneyU']
TEST_LABELS: dict[str, str] = {'KS': 'KS', 'CVM': 'CvM', 'MannWhitneyU': 'MWU'}

# Averaging modes in column order; the trimean is the reference every other mode is tested against.
MEAN_TYPES: list[str] = ['mean_0', 'trim_mean_0.05', 'trim_mean_0.1', 'tri_mean_0']
MEAN_TYPE_LABELS: dict[str, str] = {
    'mean_0': 'Mean',
    'trim_mean_0.05': 'Trimmed Mean 5%',
    'trim_mean_0.1': 'Trimmed Mean 10%',
    'tri_mean_0': 'Trimean',
}
REFERENCE: str = 'tri_mean_0'

# Number formats keep the cells numeric while appending the significance marker.
FORMAT_PLAIN: str = '0.000'
FORMAT_TRIMEAN_BETTER: str = '0.000"*"'
FORMAT_TRIMEAN_WORSE: str = '0.000"†"'

FONT: Font = Font(name='Aptos Narrow', size=11)
BOLD: Font = Font(name='Aptos Narrow', size=11, bold=True)
CENTER: Alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN: Side = Side(style='thin')


def paired_wilcoxon_p_value(x: np.ndarray, y: np.ndarray) -> float:
    """
    Same test as create_statistical_test_ranking_table.py: exact paired Wilcoxon signed-rank test.
    Identical PR AUCs (all differences zero) and NaN p-values count as p = 1.
    """
    if not np.any(x - y):
        return 1.0

    p_value: float = stats.wilcoxon(x, y, zero_method='wilcox', method=stats.PermutationMethod()).pvalue

    return 1.0 if np.isnan(p_value) else p_value


def outcome_vs_reference(reference: np.ndarray, other: np.ndarray, adjusted_p_value: float) -> int:
    """
    +1 if the trimean is significantly better, -1 if significantly worse, 0 if not significantly different.
    The direction follows the median PR AUC; equal medians fall back to the median paired difference.
    """
    if adjusted_p_value >= alpha:
        return 0

    median_difference: float = np.median(reference) - np.median(other)
    if median_difference == 0:
        median_difference = np.median(reference - other)

    return int(np.sign(median_difference))


def donors_label(proportion: str) -> str:
    return f'{round(float(proportion) * n_donors)}/{n_donors}'


def set_cell(ws: Worksheet, row: int, column: int, value, font: Font = FONT, number_format: str | None = None):
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = font
    cell.alignment = CENTER
    if number_format is not None:
        cell.number_format = number_format


def write_median_grid(ws: Worksheet, medians: dict, outcomes: dict):
    """
    Rows: donors affected x Log2 FC. Columns: one block of four averaging modes per statistical test.
    The highest median per test and row is bold.
    """
    first_value_column: int = 3
    block_width: int = len(MEAN_TYPES)

    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    set_cell(ws, 1, 1, 'Donors Affected', BOLD)
    set_cell(ws, 1, 2, 'Log2 FC', BOLD)

    for test_index, test in enumerate(STATISTICAL_TESTS):
        block_start: int = first_value_column + test_index * block_width
        ws.merge_cells(start_row=1, start_column=block_start, end_row=1, end_column=block_start + block_width - 1)
        set_cell(ws, 1, block_start, TEST_LABELS[test], BOLD)
        for mean_index, mean_type in enumerate(MEAN_TYPES):
            set_cell(ws, 2, block_start + mean_index, MEAN_TYPE_LABELS[mean_type], BOLD)

    row: int = 3
    for proportion in PROPORTIONS:
        for case in CASES:
            set_cell(ws, row, 1, donors_label(proportion))
            set_cell(ws, row, 2, float(case.split('_')[-1]))

            for test_index, test in enumerate(STATISTICAL_TESTS):
                block_start: int = first_value_column + test_index * block_width
                best: float = max(medians[(proportion, case, test, mean_type)] for mean_type in MEAN_TYPES)

                for mean_index, mean_type in enumerate(MEAN_TYPES):
                    median: float = medians[(proportion, case, test, mean_type)]
                    outcome: int = outcomes.get((proportion, case, test, mean_type), 0)
                    number_format: str = {
                        1: FORMAT_TRIMEAN_BETTER, -1: FORMAT_TRIMEAN_WORSE, 0: FORMAT_PLAIN,
                    }[outcome]
                    set_cell(
                        ws, row, block_start + mean_index, median,
                        font=BOLD if median == best else FONT,
                        number_format=number_format,
                    )
            row += 1

    # Vertical separators between the statistical-test blocks.
    for test_index in range(len(STATISTICAL_TESTS)):
        block_start: int = first_value_column + test_index * block_width
        for grid_row in range(1, row):
            cell = ws.cell(row=grid_row, column=block_start)
            cell.border = Border(left=THIN)

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 8
    for column in range(first_value_column, first_value_column + len(STATISTICAL_TESTS) * block_width):
        ws.column_dimensions[ws.cell(row=2, column=column).column_letter].width = 10
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = 'C3'

    footnote_row: int = row + 1
    footnotes: list[str] = [
        'Median PR AUC over 10 simulated datasets per scenario; the highest median per test and scenario is bold.',
        f'* Trimean significantly higher than this averaging mode; † trimean significantly lower '
        f'(exact paired Wilcoxon signed-rank test, Benjamini-Hochberg corrected over the three comparisons '
        f'per test and scenario, adjusted p < {alpha}).',
    ]
    for offset, footnote in enumerate(footnotes):
        ws.cell(row=footnote_row + offset, column=1, value=footnote).font = FONT


def main():
    medians: dict[tuple[str, str, str, str], float] = {}
    outcomes: dict[tuple[str, str, str, str], int] = {}
    pairwise_rows: list[dict] = []

    for proportion in PROPORTIONS:
        for case in CASES:
            pr_aucs: dict[str, dict[str, np.ndarray]] = {}
            for mean_type in MEAN_TYPES:
                path_to_pkl: str = os.path.join(path_to_summary_pkls, f'{mean_type}_{case}_{proportion}.pkl')

                if not os.path.exists(path_to_pkl):
                    raise FileNotFoundError(f"File not found: {path_to_pkl}")

                with open(path_to_pkl, 'rb') as f:
                    pr_aucs[mean_type] = {test: np.asarray(values) for test, values in pickle.load(f).items()}

            for test in STATISTICAL_TESTS:
                for mean_type in MEAN_TYPES:
                    medians[(proportion, case, test, mean_type)] = float(np.median(pr_aucs[mean_type][test]))

                others: list[str] = [mean_type for mean_type in MEAN_TYPES if mean_type != REFERENCE]
                reference: np.ndarray = pr_aucs[REFERENCE][test]
                p_values: list[float] = [paired_wilcoxon_p_value(reference, pr_aucs[other][test]) for other in others]
                adjusted_p_values: np.ndarray = stats.false_discovery_control(p_values, method='bh')

                for other, p_value, adjusted_p_value in zip(others, p_values, adjusted_p_values):
                    outcome: int = outcome_vs_reference(reference, pr_aucs[other][test], adjusted_p_value)
                    outcomes[(proportion, case, test, other)] = outcome

                    pairwise_rows.append({
                        'Donors Affected': donors_label(proportion),
                        'Log2 FC': float(case.split('_')[-1]),
                        'Statistical Test': TEST_LABELS[test],
                        'Comparison': f'Trimean vs {MEAN_TYPE_LABELS[other]}',
                        'Median PR AUC – Trimean': np.median(reference),
                        'Median PR AUC – Other': np.median(pr_aucs[other][test]),
                        'Wilcoxon p': p_value,
                        'BH-adjusted p': adjusted_p_value,
                        'Outcome': outcome,
                    })

    workbook: Workbook = Workbook()
    write_median_grid(workbook.active, medians, outcomes)
    workbook.active.title = 'Median PR AUC'

    # The p-values behind every marker in the grid.
    pairwise_ws: Worksheet = workbook.create_sheet('Pairwise comparisons')
    pairwise_df: pd.DataFrame = pd.DataFrame(pairwise_rows)
    pairwise_ws.append(list(pairwise_df.columns))
    for row in pairwise_df.itertuples(index=False):
        pairwise_ws.append(list(row))

    workbook.save(path_to_comparison_table_excel)


if __name__ == "__main__":
    main()
