import os

import pickle

import numpy as np
import pandas as pd

from scipy import stats

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

import argparse

parser = argparse.ArgumentParser()
parser.add_argument(
    '--path_to_summary_pkls',
    type=str,
    required=True,
)
parser.add_argument(
    '--mean_type',
    type=str,
    default='tri_mean_0',
)
parser.add_argument(
    '--path_to_friedman_table_excel',
    type=str,
    required=True,
)
parser.add_argument(
    '--n_bootstrap',
    type=int,
    default=10000,
)
parser.add_argument(
    '--seed',
    type=int,
    default=0,
)
args = parser.parse_args()

path_to_summary_pkls: str = args.path_to_summary_pkls
mean_type: str = args.mean_type
path_to_friedman_table_excel: str = args.path_to_friedman_table_excel
n_bootstrap: int = args.n_bootstrap
seed: int = args.seed

CASES: list[str] = ['CASE_0.3', 'CASE_0.5', 'CASE_1', 'CASE_1.5']
PROPORTIONS: list[str] = ['0.1', '0.3', '0.5', '0.8', '1.0']
STATISTICAL_TESTS: list[str] = ['KS', 'CVM', 'MannWhitneyU']
TEST_LABELS: dict[str, str] = {'KS': 'KS', 'CVM': 'CvM', 'MannWhitneyU': 'MWU'}
PAIRS: list[tuple[str, str]] = [('CVM', 'KS'), ('CVM', 'MannWhitneyU'), ('KS', 'MannWhitneyU')]

FONT: Font = Font(name='Aptos Narrow', size=11)
BOLD: Font = Font(name='Aptos Narrow', size=11, bold=True)
CENTER: Alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def load_pr_aucs() -> pd.DataFrame:
    """
    One row per simulated dataset (Log2 FC x proportion x replicate) with the PR AUC of every statistical test.
    The tests are paired: they all score the same simulated dataset.
    """
    rows: list[dict] = []
    for case in CASES:
        for proportion in PROPORTIONS:
            path_to_pkl: str = os.path.join(path_to_summary_pkls, f'{mean_type}_{case}_{proportion}.pkl')

            if not os.path.exists(path_to_pkl):
                raise FileNotFoundError(f"File not found: {path_to_pkl}")

            with open(path_to_pkl, 'rb') as f:
                pr_aucs: dict[str, np.ndarray] = {
                    test: np.asarray(values) for test, values in pickle.load(f).items()
                }

            for replicate in range(len(pr_aucs[STATISTICAL_TESTS[0]])):
                rows.append({
                    'Log2 FC': case.split('_')[-1],
                    'Proportion of Affected Donors': proportion,
                    'Replicate': replicate,
                    **{test: float(pr_aucs[test][replicate]) for test in STATISTICAL_TESTS},
                })

    return pd.DataFrame(rows)


def friedman_summary(blocks: pd.DataFrame, block_label: str) -> dict:
    """
    Friedman test with the statistical tests as treatments and the rows of `blocks` as blocks. Ranks are taken
    within each block (1 = highest PR AUC, ties get the average rank).

    Also reports the Iman-Davenport F correction (Demšar 2006) and Kendall's W as effect size.
    """
    n: int = len(blocks)
    k: int = len(STATISTICAL_TESTS)

    ranks: pd.DataFrame = blocks[STATISTICAL_TESTS].rank(axis=1, ascending=False, method='average')
    q, p_value = stats.friedmanchisquare(*[blocks[test] for test in STATISTICAL_TESTS])

    f_statistic: float = (n - 1) * q / (n * (k - 1) - q)
    df_1, df_2 = k - 1, (k - 1) * (n - 1)

    return {
        'Blocks': block_label,
        'n Blocks': n,
        **{f'Mean Rank {TEST_LABELS[test]}': ranks[test].mean() for test in STATISTICAL_TESTS},
        'Friedman χ²': q,
        'df': k - 1,
        'p': p_value,
        'Iman-Davenport F': f_statistic,
        'df1': df_1,
        'df2': df_2,
        'p (Iman-Davenport)': stats.f.sf(f_statistic, df_1, df_2),
        "Kendall's W": q / (n * (k - 1)),
    }


def holm_adjust(p_values: list[float]) -> np.ndarray:
    """Holm step-down adjustment of a family of p-values."""
    p_values: np.ndarray = np.asarray(p_values)
    m: int = len(p_values)
    order: np.ndarray = np.argsort(p_values)

    adjusted: np.ndarray = np.empty(m)
    running_max: float = 0.0
    for position, index in enumerate(order):
        running_max = max(running_max, (m - position) * p_values[index])
        adjusted[index] = min(1.0, running_max)

    return adjusted


def hierarchical_bootstrap_ci(differences_per_scenario: list[np.ndarray], rng: np.random.Generator) -> np.ndarray:
    """
    95% CI of the mean paired PR AUC difference. Resamples scenarios, then replicates within each drawn scenario,
    so the CI reflects both sources of variation instead of treating all datasets as independent.
    """
    n_scenarios: int = len(differences_per_scenario)
    bootstrap_means: np.ndarray = np.empty(n_bootstrap)

    for b in range(n_bootstrap):
        drawn_scenarios: np.ndarray = rng.integers(0, n_scenarios, n_scenarios)
        bootstrap_means[b] = np.mean([
            rng.choice(differences_per_scenario[s], len(differences_per_scenario[s])).mean()
            for s in drawn_scenarios
        ])

    return np.percentile(bootstrap_means, [2.5, 97.5])


def pairwise_post_hoc(pr_aucs: pd.DataFrame, scenario_means: pd.DataFrame) -> pd.DataFrame:
    """
    Pairwise follow-up to the Friedman test: Wilcoxon signed-rank test on the scenario means (Benavoli et al. 2016),
    Holm-adjusted over the three pairs, with the mean PR AUC difference and its hierarchical bootstrap CI.
    """
    rng: np.random.Generator = np.random.default_rng(seed)
    rows: list[dict] = []

    for a, b in PAIRS:
        scenario_differences: pd.Series = scenario_means[a] - scenario_means[b]
        dataset_differences: pd.Series = pr_aucs[a] - pr_aucs[b]
        differences_per_scenario: list[np.ndarray] = [
            group.to_numpy()
            for _, group in dataset_differences.groupby([pr_aucs['Log2 FC'], pr_aucs['Proportion of Affected Donors']])
        ]
        ci_low, ci_high = hierarchical_bootstrap_ci(differences_per_scenario, rng)

        rows.append({
            'Comparison': f'{TEST_LABELS[a]} vs {TEST_LABELS[b]}',
            'Mean ΔPR AUC': dataset_differences.mean(),
            'Median ΔPR AUC': dataset_differences.median(),
            '95% CI Low': ci_low,
            '95% CI High': ci_high,
            'Scenarios Won by First': int((scenario_differences > 0).sum()),
            'Scenarios Won by Second': int((scenario_differences < 0).sum()),
            'Datasets Won by First (%)': 100 * (dataset_differences > 0).mean(),
            'Wilcoxon p': stats.wilcoxon(scenario_means[a], scenario_means[b], zero_method='wilcox').pvalue,
        })

    post_hoc_df: pd.DataFrame = pd.DataFrame(rows)
    post_hoc_df.insert(len(post_hoc_df.columns), 'Holm-adjusted p', holm_adjust(post_hoc_df['Wilcoxon p']))

    return post_hoc_df


def set_cell(ws: Worksheet, row: int, column: int, value, font: Font = FONT, number_format: str | None = None):
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = font
    cell.alignment = CENTER
    if number_format is not None:
        cell.number_format = number_format


def number_format_for(column: str, value) -> str | None:
    if not isinstance(value, float):
        return None
    if column == 'p' or column.endswith(' p') or column.startswith('p ('):
        return '0.00E+00'
    if 'Rank' in column or '%' in column or column in ['Friedman χ²', 'Iman-Davenport F', "Kendall's W"]:
        return '0.00'
    return '0.000'


def write_dataframe(ws: Worksheet, df: pd.DataFrame, start_row: int = 1, column_width: float = 13.0) -> int:
    """Writes `df` with a bold header at `start_row`; returns the first empty row below it."""
    for column_index, column in enumerate(df.columns, start=1):
        set_cell(ws, start_row, column_index, column, font=BOLD)
        ws.column_dimensions[ws.cell(row=start_row, column=column_index).column_letter].width = column_width

    for row_offset, row in enumerate(df.itertuples(index=False), start=1):
        for column_index, (column, value) in enumerate(zip(df.columns, row), start=1):
            value = value.item() if isinstance(value, np.generic) else value
            set_cell(ws, start_row + row_offset, column_index, value, number_format=number_format_for(column, value))

    return start_row + len(df) + 1


def write_notes(ws: Worksheet, start_row: int, notes: list[str]):
    for offset, note in enumerate(notes):
        cell = ws.cell(row=start_row + offset, column=1, value=note)
        cell.font = FONT


def main():
    pr_aucs: pd.DataFrame = load_pr_aucs()

    # One block per simulation scenario: the mean PR AUC over its replicate datasets.
    scenario_means: pd.DataFrame = (
        pr_aucs.groupby(['Log2 FC', 'Proportion of Affected Donors'], sort=False)[STATISTICAL_TESTS]
        .mean()
        .reset_index()
    )

    friedman_df: pd.DataFrame = pd.DataFrame([
        friedman_summary(scenario_means, 'Scenarios (mean over replicates)'),
        friedman_summary(pr_aucs, 'Datasets (replicates as blocks)'),
    ])
    post_hoc_df: pd.DataFrame = pairwise_post_hoc(pr_aucs, scenario_means)

    ranks: pd.DataFrame = scenario_means[STATISTICAL_TESTS].rank(axis=1, ascending=False, method='average')
    ranks_df: pd.DataFrame = pd.concat([
        scenario_means[['Log2 FC', 'Proportion of Affected Donors']],
        scenario_means[STATISTICAL_TESTS].rename(columns=lambda test: f'Mean PR AUC {TEST_LABELS[test]}'),
        ranks.rename(columns=lambda test: f'Rank {TEST_LABELS[test]}'),
    ], axis=1)
    rank_sum_row: dict = {'Log2 FC': 'Rank Sum', **{f'Rank {TEST_LABELS[t]}': ranks[t].sum() for t in STATISTICAL_TESTS}}
    mean_rank_row: dict = {'Log2 FC': 'Mean Rank', **{f'Rank {TEST_LABELS[t]}': ranks[t].mean() for t in STATISTICAL_TESTS}}
    ranks_df = pd.concat([ranks_df, pd.DataFrame([rank_sum_row, mean_rank_row])], ignore_index=True)
    ranks_df = ranks_df.astype(object).where(ranks_df.notna(), None)

    workbook: Workbook = Workbook()

    friedman_ws: Worksheet = workbook.active
    friedman_ws.title = 'Friedman test'
    next_row: int = write_dataframe(friedman_ws, friedman_df)
    write_notes(friedman_ws, next_row + 1, [
        f'PR AUCs of the {mean_type} runs. Treatments: KS, CvM, MWU. Rank 1 = highest PR AUC; ties get the average rank.',
        'Primary analysis: the 20 simulation scenarios (Log2 FC x proportion) as blocks, each the mean over its '
        'replicate datasets.',
        'The dataset-level row treats every replicate as an independent block and therefore overstates the evidence; '
        'it is listed for reference only.',
        'Iman-Davenport F: less conservative F-approximation of the Friedman statistic (Demšar 2006). '
        "Kendall's W = χ² / (n (k - 1)): 0 = no agreement between blocks on the ranking, 1 = identical ranking.",
    ])

    post_hoc_ws: Worksheet = workbook.create_sheet('Pairwise post hoc')
    next_row = write_dataframe(post_hoc_ws, post_hoc_df)
    write_notes(post_hoc_ws, next_row + 1, [
        'Wilcoxon signed-rank test on the 20 scenario means, Holm-adjusted over the three pairs (Benavoli et al. 2016).',
        'ΔPR AUC = first minus second test, over all datasets. 95% CI from a hierarchical bootstrap '
        f'(scenarios, then replicates within scenarios; {n_bootstrap} resamples, seed {seed}).',
        'Scenarios won: scenario-mean PR AUC higher. Datasets won: share of datasets with a higher PR AUC.',
    ])

    ranks_ws: Worksheet = workbook.create_sheet('Ranks per scenario')
    write_dataframe(ranks_ws, ranks_df)

    data_ws: Worksheet = workbook.create_sheet('PR AUCs per dataset')
    write_dataframe(data_ws, pr_aucs.rename(columns={test: f'PR AUC {TEST_LABELS[test]}' for test in STATISTICAL_TESTS}))

    workbook.save(path_to_friedman_table_excel)


if __name__ == "__main__":
    main()
